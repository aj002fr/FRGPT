"""Parse weekly Excel calendar into normalized scheduled-event records."""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from .models import ScheduledEvent


logger = logging.getLogger(__name__)
ET_TZ = ZoneInfo("America/New_York")

DATE_ROW_RE = re.compile(
    r"^(Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday)\s+"
    r"([A-Za-z]+)\s+(\d{1,2})\s+(\d{4})"
)
HOLIDAY_DATE_RE = re.compile(r"^([A-Za-z]{3})\s+(\d{1,2})")

MONTH_LOOKUP = {
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}


@dataclass(frozen=True)
class ParseMetrics:
    rows_total: int = 0
    rows_parsed: int = 0
    rows_skipped_invalid: int = 0
    rows_past: int = 0


def _parse_date_row(value: object) -> date | None:
    if not value:
        return None
    text = str(value).strip()
    match = DATE_ROW_RE.match(text)
    if not match:
        return None
    month_name = match.group(2)
    day_str = match.group(3)
    year_str = match.group(4)
    return datetime.strptime(f"{month_name} {day_str} {year_str}", "%B %d %Y").date()


def _parse_time_cell(value: object) -> time | None:
    """Parse first timestamp from Excel time cell."""
    if value is None:
        return None
    if isinstance(value, time):
        return value.replace(second=value.second or 0, microsecond=0)
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)

    text = str(value).strip()
    if not text:
        return None

    # Handle range style like "3:15 - 4 AM" by taking earliest component.
    if "-" in text:
        left, right = [part.strip() for part in text.split("-", 1)]
        right_suffix = ""
        if re.search(r"\bAM\b", right, re.IGNORECASE):
            right_suffix = " AM"
        elif re.search(r"\bPM\b", right, re.IGNORECASE):
            right_suffix = " PM"
        if not re.search(r"\b(AM|PM)\b", left, re.IGNORECASE) and right_suffix:
            text = f"{left}{right_suffix}"
        else:
            text = left

    for fmt in ("%H:%M:%S", "%H:%M", "%I:%M %p", "%I %p"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    return None


def _normalize_country(value: object) -> str:
    if value is None:
        return "UNK"
    raw = str(value).strip()
    if not raw:
        return "UNK"
    return raw.upper()


def _clean_event_name(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\u00a0", " ").replace("\t", " ").strip()
    return re.sub(r"\s+", " ", text)


def _parse_holiday_sheet(workbook, sheet_name: str) -> dict[str, set[date]]:
    """Parse region holiday table into country->dates map."""
    if sheet_name not in workbook.sheetnames:
        return {}

    ws = workbook[sheet_name]
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 300), values_only=True))
    if len(rows) < 2:
        return {}

    # Header row typically contains "Date (2026)" + region columns.
    header_row = rows[1]
    region_columns: dict[int, str] = {}
    for idx, header in enumerate(header_row):
        if idx == 0 or not header:
            continue
        label = str(header).lower()
        if "japan" in label:
            region_columns[idx] = "JP"
        elif "uk" in label or "england" in label:
            region_columns[idx] = "GB"
        elif "usa" in label or "federal" in label:
            region_columns[idx] = "US"
        elif "europe" in label or "eu " in label:
            region_columns[idx] = "EU"

    # Infer default year from first header cell "Date (2026)".
    default_year = datetime.now(ET_TZ).year
    date_header = str(header_row[0]) if header_row and header_row[0] else ""
    year_match = re.search(r"(\d{4})", date_header)
    if year_match:
        default_year = int(year_match.group(1))

    result: dict[str, set[date]] = {}
    for row in rows[2:]:
        if not row or not row[0]:
            continue
        date_text = str(row[0]).strip()
        date_match = HOLIDAY_DATE_RE.match(date_text)
        if not date_match:
            continue
        month = MONTH_LOOKUP.get(date_match.group(1).lower())
        if not month:
            continue
        day_num = int(date_match.group(2))
        try:
            row_date = date(default_year, month, day_num)
        except ValueError:
            continue

        for idx, region in region_columns.items():
            if idx >= len(row):
                continue
            marker = row[idx]
            marker_text = "" if marker is None else str(marker).strip()
            if marker_text and marker_text not in {"-", "—", "–", "•"}:
                result.setdefault(region, set()).add(row_date)
    return result


def _detect_base_col(ws) -> int:
    """Detect the leading column index that holds the date/time values.

    Two layouts exist in the wild:
      - offset layout: column A empty, B=date/time, C=country, D=event_name
      - flush layout: column A=Time, B=Ctry, C=Event (no leading empty column)

    We scan the first rows and pick the column where the weekday date headers
    ("Monday June 29 2026") appear most often.
    """
    offset0_hits = 0
    offset1_hits = 0
    scanned = 0
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 80), values_only=True):
        scanned += 1
        col0 = row[0] if len(row) > 0 else None
        col1 = row[1] if len(row) > 1 else None
        if _parse_date_row(col0):
            offset0_hits += 1
        if _parse_date_row(col1):
            offset1_hits += 1
    # Prefer the column with date-header matches; default to the historical
    # offset layout (base col 1) when neither produced a hit.
    if offset0_hits > offset1_hits:
        return 0
    if offset1_hits > 0:
        return 1
    return 1


def parse_excel_calendar(
    excel_path: str | Path,
) -> tuple[list[ScheduledEvent], dict[str, set[date]], ParseMetrics]:
    """
    Parse Calendar sheet to normalized ScheduledEvent rows for DB upsert.

    Returns:
      (events, holiday_dates_by_country, metrics)
    """
    source_path = Path(excel_path)
    source_file = source_path.name
    wb = load_workbook(filename=str(source_path), data_only=True, read_only=True)
    if "Calendar" not in wb.sheetnames:
        raise ValueError(f"Calendar sheet not found in {source_path}")

    ws = wb["Calendar"]
    events: list[ScheduledEvent] = []
    metrics = ParseMetrics()
    current_date: date | None = None
    now_et = datetime.now(ET_TZ)

    base_col = _detect_base_col(ws)

    rows_total = 0
    rows_parsed = 0
    rows_invalid = 0
    rows_past = 0

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        rows_total += 1

        # Columns are relative to the detected base column:
        # base: date/time, base+1: country, base+2: event_name.
        col_b = row[base_col] if len(row) > base_col else None  # date or time
        col_c = row[base_col + 1] if len(row) > base_col + 1 else None  # country
        col_d = row[base_col + 2] if len(row) > base_col + 2 else None  # event_name

        parsed_date = _parse_date_row(col_b)
        if parsed_date:
            current_date = parsed_date
            continue

        if current_date is None:
            continue

        event_name = _clean_event_name(col_d)
        if not event_name:
            continue

        event_time = _parse_time_cell(col_b)
        if event_time is None:
            rows_invalid += 1
            continue

        country = _normalize_country(col_c)
        scheduled_dt = datetime.combine(current_date, event_time, tzinfo=ET_TZ)

        if scheduled_dt < now_et:
            rows_past += 1

        event_id = ScheduledEvent.build_event_id(event_name, scheduled_dt, source_file)
        events.append(
            ScheduledEvent(
                event_id=event_id,
                event_name=event_name,
                country=country,
                scheduled_time_et=scheduled_dt,
                source_file=source_file,
                is_active=True,
            )
        )
        rows_parsed += 1

    holidays = {}
    for holiday_sheet in ("holidays", "holiday"):
        parsed = _parse_holiday_sheet(wb, holiday_sheet)
        for region, dates in parsed.items():
            holidays.setdefault(region, set()).update(dates)

    metrics = ParseMetrics(
        rows_total=rows_total,
        rows_parsed=rows_parsed,
        rows_skipped_invalid=rows_invalid,
        rows_past=rows_past,
    )
    logger.info(
        "Parsed Excel calendar: total=%s parsed=%s invalid=%s past=%s file=%s",
        metrics.rows_total,
        metrics.rows_parsed,
        metrics.rows_skipped_invalid,
        metrics.rows_past,
        source_file,
    )
    return events, holidays, metrics

