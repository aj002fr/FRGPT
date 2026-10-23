#!/usr/bin/env python3
"""Export the 'By surprise value' table for a series to a clean, formatted Excel.

Usage: python scripts/export_cpi_core_surprise_table.py --series core_ppi
"""
from __future__ import annotations

import argparse
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).parent.parent

SERIES_META = {
    "core_cpi": ("Core CPI MoM", "core_cpi_surprise_buckets"),
    "headline_cpi": ("Headline CPI MoM", "headline_cpi_surprise_buckets"),
    "core_ppi": ("Core PPI MoM", "core_ppi_surprise_buckets"),
    "headline_ppi": ("Headline PPI MoM", "headline_ppi_surprise_buckets"),
}

COLS = [
    ("surprise_pp", "Surprise (pp)"),
    ("mean_std_surprise_z", "z"),
    ("n", "n"),
    ("pctl_of_distribution", "Pctl (%)"),
    ("expected_abs_move_bps", "Expected |move| (bps)"),
    ("median_abs_move_bps", "Median |move| (bps)"),
    ("mean_signed_move_bps", "Mean signed (bps)"),
    ("max_abs_move_bps", "Max |move| (bps)"),
    ("avg_max_excursion_bps", "Avg max excursion (bps)"),
    ("max_excursion_bps", "Max excursion (bps)"),
    ("max_excursion_date", "Max excursion date"),
]

def build_notes(label: str) -> list[str]:
    return [
        f"{label} surprise -> TY (10Y futures) reaction. Sample: monthly releases Jan-2022 .. May-2026.",
        "Move convention: 8:30am -> 5pm ET close-to-close in bps (TY DV01 conversion via points_to_bps).",
        "Negative signed = TY price down / yields up (hawkish reaction to hot prints).",
        "'Max excursion' = one-sided intraday max move from the 8:30am open within the post window.",
        "z = surprise / consensus std dev. Pctl = share of historical surprises strictly below this value.",
        "Surprises are quoted to ~0.1pp granularity, so values are discrete.",
        "Thin tails (n=1) should be treated as anecdotes, not robust estimates.",
        "PPI is historically a weaker Treasury mover than CPI; signed reactions are noisier.",
        "Window limited to 2022+ because TY 1-min data in PostgreSQL starts there.",
    ]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--series", default="core_cpi", choices=list(SERIES_META.keys()))
    args = ap.parse_args()
    label, folder = SERIES_META[args.series]
    base = PROJECT_ROOT / "workspace" / "analysis" / folder
    src_db = base / f"{folder}.db"
    out_xlsx = base / f"{args.series}_surprise_value_table.xlsx"
    notes = build_notes(label)

    with sqlite3.connect(str(src_db)) as conn:
        df = pd.read_sql_query("SELECT * FROM by_surprise_value ORDER BY surprise_pp", conn)

    out = df[[c for c, _ in COLS]].rename(columns=dict(COLS))

    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as xl:
        out.to_excel(xl, sheet_name="By Surprise Value", index=False, startrow=1)
        ws = xl.sheets["By Surprise Value"]

        title = f"{label} Surprise -> TY Move (bps)"
        ws.cell(row=1, column=1, value=title)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLS))
        ws.cell(row=1, column=1).font = Font(bold=True, size=13)
        ws.cell(row=1, column=1).alignment = Alignment(horizontal="left")

        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(bold=True, color="FFFFFF")
        thin = Side(style="thin", color="BBBBBB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)

        header_row = 2
        for j in range(len(COLS)):
            c = ws.cell(row=header_row, column=j + 1)
            c.fill = header_fill
            c.font = header_font
            c.alignment = center
            c.border = border

        n_rows = len(out)
        hot_fill = PatternFill("solid", fgColor="FBE5E1")   # hawkish (positive surprise)
        cool_fill = PatternFill("solid", fgColor="E2F0E2")  # dovish (negative surprise)
        for i in range(n_rows):
            excel_row = header_row + 1 + i
            surprise = out.iloc[i]["Surprise (pp)"]
            row_fill = hot_fill if surprise > 0 else (cool_fill if surprise < 0 else None)
            for j in range(len(COLS)):
                c = ws.cell(row=excel_row, column=j + 1)
                c.border = border
                c.alignment = Alignment(horizontal="center")
                if row_fill is not None:
                    c.fill = row_fill

        widths = [13, 7, 6, 9, 17, 16, 15, 14, 19, 17, 18]
        for j, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(j + 1)].width = w
        ws.freeze_panes = "A3"

        notes_df = pd.DataFrame({"Notes / caveats": notes})
        notes_df.to_excel(xl, sheet_name="Notes", index=False)
        nws = xl.sheets["Notes"]
        nws.column_dimensions["A"].width = 110
        nws.cell(row=1, column=1).font = Font(bold=True)
        for r in range(2, len(notes) + 2):
            nws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")

    print(f"Wrote: {out_xlsx}")
    print(out.to_string(index=False))


if __name__ == "__main__":
    main()
